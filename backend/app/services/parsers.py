from pathlib import Path
import pdfplumber
from docx import Document as DocxDocument
from pptx import Presentation
from openpyxl import load_workbook

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx"}


def extract_text(path: Path) -> str:
    """Extract readable text while retaining page/sheet context for citations."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        with pdfplumber.open(path) as pdf:
            return "\n".join(f"[Page {i}]\n{page.extract_text() or ''}" for i, page in enumerate(pdf.pages, 1))
    if suffix == ".docx":
        return "\n".join(p.text for p in DocxDocument(path).paragraphs if p.text.strip())
    if suffix == ".pptx":
        presentation = Presentation(path)
        return "\n".join(f"[Slide {i}]\n" + "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text")) for i, slide in enumerate(presentation.slides, 1))
    if suffix == ".xlsx":
        book = load_workbook(path, data_only=True, read_only=True)
        return "\n".join(f"[Sheet {sheet.title}]\n" + "\n".join(" | ".join(str(cell) for cell in row if cell is not None) for row in sheet.iter_rows(values_only=True)) for sheet in book.worksheets)
    raise ValueError(f"Unsupported file type: {suffix}")
